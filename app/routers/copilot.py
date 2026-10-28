from fastapi import APIRouter, Depends, Request, UploadFile, File, HTTPException
from sqlalchemy.orm import Session
from pydantic import BaseModel, Field
from typing import Any, List, Literal, Optional
from datetime import datetime, timezone
import logging
import re
import json
from ..database import get_db
from ..copilot_service import (
    process_copilot_chat,
    cleanup_old_chats,
    extract_signed_report_artifact,
    extract_signed_sql_export,
)
from ..models import Usuario, CopilotChat, CopilotMessage
from ..api.deps import get_current_user
from ..core.rate_limit import limiter
from .. import models

router = APIRouter(prefix="/copilot", tags=["copilot"])
logger = logging.getLogger(__name__)
MAX_CONTEXT_FILE_SIZE = 1_048_576
MAX_CONTEXT_CHARS = 200_000
MAX_CHAT_CONTEXT_MESSAGES = 20
MAX_CHAT_HISTORY_MESSAGES = 100
MAX_CHAT_LIST = 100
MAX_LIBRERIA_DOCS_PER_QUERY = 10
MAX_DOCUMENT_PAGES = 50


def _filas_informe(valor: Any) -> list[dict[str, Any]]:
    if isinstance(valor, list):
        return [fila for fila in valor if isinstance(fila, dict)]
    if isinstance(valor, dict):
        return [valor]
    return []


def _nombre_hoja_informe(nombre: str, existentes: set[str]) -> str:
    """Devuelve un nombre de hoja Excel estable y válido (máximo 31 caracteres)."""
    base = re.sub(r"[\\/*?:\[\]]", "_", nombre.replace("_", " ").strip().title())[:31] or "Evidencia"
    candidato, indice = base, 2
    while candidato in existentes:
        sufijo = f" {indice}"
        candidato = f"{base[:31 - len(sufijo)]}{sufijo}"
        indice += 1
    existentes.add(candidato)
    return candidato


def _numero_informe(valor: Any) -> float:
    try:
        return float(valor or 0)
    except (TypeError, ValueError):
        return 0.0


def _eur_informe(valor: Any) -> str:
    return f"€{_numero_informe(valor):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _kpis_informe(informe: dict[str, Any]) -> list[dict[str, str]]:
    """Resume métricas ya calculadas, sin volver a derivarlas desde la prosa."""
    resultados = informe.get("resultados") if isinstance(informe.get("resultados"), dict) else {}
    resumen = _filas_informe(resultados.get("resumen"))
    fila = resumen[0] if resumen else {}
    ventas_actuales = fila.get("ventas_periodo_actual")
    ventas_anteriores = fila.get("ventas_periodo_anterior")
    variacion = fila.get("ventas_variacion_absoluta")
    variacion_pct = fila.get("ventas_variacion_pct")
    metricas = []
    if ventas_actuales is not None:
        metricas.append({"Métrica": "Ventas actuales", "Valor": _eur_informe(ventas_actuales)})
    if ventas_anteriores is not None:
        metricas.append({"Métrica": "Ventas comparables", "Valor": _eur_informe(ventas_anteriores)})
    if variacion is not None:
        texto_pct = f" ({_numero_informe(variacion_pct):.2f} %)".replace(".", ",") if variacion_pct is not None else ""
        metricas.append({"Métrica": "Variación de ventas", "Valor": f"{_eur_informe(variacion)}{texto_pct}"})
    for clave, etiqueta in (("margen_periodo_actual", "Margen bruto actual"), ("mgd_periodo_actual", "MGD actual")):
        if fila.get(clave) is not None:
            metricas.append({"Métrica": etiqueta, "Valor": _eur_informe(fila[clave])})
    return metricas


def _exportar_informe_ejecutivo(informe: dict[str, Any], message_id: int, formato: str):
    """Construye un artefacto reproducible desde el evidence bundle firmado."""
    import io
    from fastapi.responses import StreamingResponse

    if formato == "xlsx":
        import pandas as pd
        from openpyxl.styles import Alignment, Font, PatternFill

        salida = io.BytesIO()
        resultados = informe.get("resultados") if isinstance(informe.get("resultados"), dict) else {}
        with pd.ExcelWriter(salida, engine="openpyxl") as writer:
            resumen = pd.DataFrame([
                {"Campo": "Tipo", "Valor": "Informe ejecutivo verificable"},
                {"Campo": "Período actual", "Valor": f"{informe.get('periodo_actual', {}).get('inicio', 'N/D')} a {informe.get('periodo_actual', {}).get('fin', 'N/D')}"},
                {"Campo": "Período comparable", "Valor": f"{informe.get('periodo_anterior', {}).get('inicio', 'N/D')} a {informe.get('periodo_anterior', {}).get('fin', 'N/D')}"},
                {"Campo": "Trazabilidad", "Valor": f"Mensaje Copilot {message_id}"},
            ])
            resumen.to_excel(writer, sheet_name="Resumen", index=False)
            kpis = _kpis_informe(informe)
            if kpis:
                pd.DataFrame(kpis).to_excel(writer, sheet_name="Resumen", index=False, startrow=len(resumen) + 3)
            hojas = {"Resumen"}
            for nombre, valor in resultados.items():
                filas = _filas_informe(valor)
                if filas:
                    pd.DataFrame(filas).to_excel(writer, sheet_name=_nombre_hoja_informe(str(nombre), hojas), index=False)
            for worksheet in writer.book.worksheets:
                for cell in worksheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="0071E3")
                for columna in worksheet.columns:
                    ancho = min(max(len(str(cell.value or "")) for cell in columna) + 2, 42)
                    worksheet.column_dimensions[columna[0].column_letter].width = ancho
                for fila in worksheet.iter_rows():
                    for cell in fila:
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
        salida.seek(0)
        return StreamingResponse(
            salida,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=informe_copilot_{message_id}.xlsx"},
        )

    if formato == "pdf":
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import cm
            from reportlab.graphics.charts.barcharts import VerticalBarChart
            from reportlab.graphics.shapes import Drawing
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
            from xml.sax.saxutils import escape
        except ImportError as error:
            raise HTTPException(
                status_code=503,
                detail="La generación de PDF no está disponible en este servidor. Descarga el informe en Excel o instala las dependencias declaradas.",
            ) from error

        salida = io.BytesIO()
        documento = SimpleDocTemplate(salida, pagesize=landscape(A4), rightMargin=1.1 * cm, leftMargin=1.1 * cm, topMargin=1.1 * cm, bottomMargin=1.1 * cm)
        estilos = getSampleStyleSheet()
        contenido = [
            Paragraph("Informe ejecutivo verificable", estilos["Title"]),
            Paragraph(
                f"Período actual: {escape(str(informe.get('periodo_actual', {}).get('inicio', 'N/D')))} a {escape(str(informe.get('periodo_actual', {}).get('fin', 'N/D')))}<br/>"
                f"Período comparable: {escape(str(informe.get('periodo_anterior', {}).get('inicio', 'N/D')))} a {escape(str(informe.get('periodo_anterior', {}).get('fin', 'N/D')))}<br/>"
                f"Trazabilidad: mensaje Copilot {message_id}. Datos generados desde evidencia SQL validada.",
                estilos["BodyText"],
            ),
            Spacer(1, 0.35 * cm),
        ]
        resultados = informe.get("resultados") if isinstance(informe.get("resultados"), dict) else {}
        kpis = _kpis_informe(informe)
        if kpis:
            contenido.append(Paragraph("Resumen ejecutivo", estilos["Heading2"]))
            tabla_kpis = Table([[item["Métrica"] for item in kpis], [item["Valor"] for item in kpis]], hAlign="LEFT")
            tabla_kpis.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#10243E")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#EAF3FF")),
                ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#10243E")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D2D2D7")),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]))
            contenido.extend((tabla_kpis, Spacer(1, 0.25 * cm)))

        resumen_evidencia = _filas_informe(resultados.get("resumen"))
        if resumen_evidencia:
            fila_resumen = resumen_evidencia[0]
            comparable = _numero_informe(fila_resumen.get("ventas_periodo_anterior"))
            actual = _numero_informe(fila_resumen.get("ventas_periodo_actual"))
            if fila_resumen.get("ventas_periodo_anterior") is not None and fila_resumen.get("ventas_periodo_actual") is not None and comparable >= 0 and actual >= 0:
                contenido.append(Paragraph("Ventas comparables frente a actuales", estilos["Heading2"]))
                dibujo = Drawing(420, 170)
                barras = VerticalBarChart()
                barras.x, barras.y, barras.width, barras.height = 45, 30, 330, 110
                barras.data = [[comparable, actual]]
                barras.categoryAxis.categoryNames = ["Comparable", "Actual"]
                barras.valueAxis.valueMin = 0
                barras.valueAxis.labels.fontSize = 7
                barras.categoryAxis.labels.fontSize = 8
                barras.bars[0].fillColor = colors.HexColor("#0071E3")
                dibujo.add(barras)
                contenido.extend((dibujo, Spacer(1, 0.25 * cm)))
        for nombre, valor in resultados.items():
            filas = _filas_informe(valor)
            if not filas:
                continue
            contenido.append(Paragraph(escape(str(nombre).replace("_", " ").title()), estilos["Heading2"]))
            columnas = list(filas[0].keys())[:7]
            tabla = [[Paragraph(escape(str(columna)), estilos["BodyText"]) for columna in columnas]]
            for fila in filas[:20]:
                tabla.append([Paragraph(escape(str(fila.get(columna) if fila.get(columna) is not None else "N/D"))[:150], estilos["BodyText"]) for columna in columnas])
            ancho_total = 27.0 * cm
            componente = Table(tabla, repeatRows=1, hAlign="LEFT", colWidths=[ancho_total / max(len(columnas), 1)] * len(columnas))
            componente.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0071E3")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D2D2D7")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F7")]),
            ]))
            contenido.extend((componente, Spacer(1, 0.3 * cm)))
            if len(filas) > 20:
                contenido.append(Paragraph(f"Se incluyen 20 de {len(filas)} filas en este bloque. El Excel conserva todas las filas.", estilos["Italic"]))
        documento.build(contenido)
        salida.seek(0)
        return StreamingResponse(
            salida,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename=informe_copilot_{message_id}.pdf"},
        )

    raise HTTPException(status_code=400, detail="El informe ejecutivo se puede descargar en PDF o XLSX.")


def _fecha_mensaje_api(fecha: datetime) -> str:
    """Expone los timestamps históricos como UTC inequívoco para el navegador."""
    if fecha.tzinfo is None:
        fecha = fecha.replace(tzinfo=timezone.utc)
    return fecha.isoformat()

class ChatMessage(BaseModel):
    role: Literal["user"]
    content: str = Field(..., max_length=2000)

class ChatRequest(BaseModel):
    chat_id: Optional[int] = None
    history: List[ChatMessage] = Field(..., min_length=1, max_length=MAX_CHAT_CONTEXT_MESSAGES)
    # Se aceptan valores heredados mientras los clientes con caché se actualizan;
    # el servicio los normaliza a Luna, Terra o Sol antes de llamar a OpenAI.
    model_preference: Literal["luna", "terra", "sol", "fast", "thinking", "ultra_thinking"] = "terra"
    libreria_doc_ids: Optional[List[int]] = Field(default=None, max_length=MAX_LIBRERIA_DOCS_PER_QUERY)

class ChatResponse(BaseModel):
    reply: str
    chat_id: int
    message_id: int

class RenameChatRequest(BaseModel):
    titulo: str = Field(..., min_length=1, max_length=100)

class ContextoRequest(BaseModel):
    contexto_negocio: str = Field(..., max_length=MAX_CONTEXT_CHARS)


class CopilotCapabilitiesResponse(BaseModel):
    inventario_disponible: bool
    ventas_disponibles: bool
    abc_ventas_disponible: bool

@router.get("/context")
def get_context(db: Session = Depends(get_db), current_user: Usuario = Depends(get_current_user)):
    empresa = db.query(models.Empresa).filter(models.Empresa.id == current_user.empresa_id).first()
    if not empresa:
        return {"contexto_negocio": ""}
    return {"contexto_negocio": empresa.contexto_negocio or ""}


@router.get("/capabilities", response_model=CopilotCapabilitiesResponse)
def get_capabilities(db: Session = Depends(get_db), current_user: Usuario = Depends(get_current_user)):
    """Expone solo la disponibilidad de datos necesaria para orientar el Copilot."""
    inventario_disponible = db.query(models.InventarioSnapshot.producto_id).join(
        models.Producto, models.Producto.id == models.InventarioSnapshot.producto_id
    ).filter(models.Producto.empresa_id == current_user.empresa_id).first() is not None
    ventas_disponibles = db.query(models.VentaHistorica.id).join(
        models.Producto, models.Producto.id == models.VentaHistorica.producto_id
    ).filter(models.Producto.empresa_id == current_user.empresa_id).first() is not None
    return {
        "inventario_disponible": inventario_disponible,
        "ventas_disponibles": ventas_disponibles,
        "abc_ventas_disponible": ventas_disponibles,
    }

@router.put("/context")
def update_context(payload: ContextoRequest, db: Session = Depends(get_db), current_user: Usuario = Depends(get_current_user)):
    empresa = db.query(models.Empresa).filter(models.Empresa.id == current_user.empresa_id).first()
    if not empresa:
        raise HTTPException(status_code=404, detail="Empresa no encontrada")
    empresa.contexto_negocio = payload.contexto_negocio.strip()
    db.commit()
    return {"success": True}

@router.post("/context/upload")
async def upload_context_document(file: UploadFile = File(...), db: Session = Depends(get_db), current_user: Usuario = Depends(get_current_user)):
    raw_filename = file.filename or "documento"
    if file.size and file.size > MAX_CONTEXT_FILE_SIZE:
        raise HTTPException(status_code=400, detail="El archivo excede el tamaño máximo permitido de 1MB.")
        
    content = await file.read(MAX_CONTEXT_FILE_SIZE + 1)
    if len(content) > MAX_CONTEXT_FILE_SIZE:
        raise HTTPException(status_code=400, detail="El archivo excede el tamaño máximo permitido de 1MB.")
        
    filename = raw_filename.lower()
    text_content = ""
    
    try:
        if filename.endswith(".txt") or filename.endswith(".csv"):
            text_content = content.decode("utf-8")
        elif filename.endswith(".pdf"):
            from pypdf import PdfReader
            import io
            reader = PdfReader(io.BytesIO(content))
            if len(reader.pages) > MAX_DOCUMENT_PAGES:
                raise HTTPException(status_code=400, detail="El documento supera el máximo de 50 páginas.")
            for page in reader.pages:
                text_content += (page.extract_text() or "") + "\n"
        elif filename.endswith(".docx"):
            from docx import Document
            import io
            doc = Document(io.BytesIO(content))
            for para in doc.paragraphs:
                text_content += para.text + "\n"
        else:
            raise HTTPException(status_code=400, detail="Formato de archivo no soportado. Usa .txt, .csv, .pdf o .docx")
            
        new_context = ""
        empresa = db.query(models.Empresa).filter(models.Empresa.id == current_user.empresa_id).first()
        if empresa:
            current_context = empresa.contexto_negocio or ""
            new_context = current_context + f"\n\n--- Contenido de {raw_filename} ---\n{text_content.strip()}"
            if len(new_context) > MAX_CONTEXT_CHARS:
                raise HTTPException(status_code=400, detail="El contexto del negocio no puede superar 200.000 caracteres.")
            empresa.contexto_negocio = new_context
            db.commit()
            
        return {"success": True, "extracted_text": text_content, "full_context": new_context}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"[COPILOT UPLOAD] Error processing file {raw_filename}: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail="Error procesando el archivo. Comprueba que el formato sea correcto.")

@router.get("/chats")
def get_chats(db: Session = Depends(get_db), current_user: Usuario = Depends(get_current_user)):
    cleanup_old_chats(db, current_user.id)
    chats = db.query(CopilotChat).filter(
        CopilotChat.usuario_id == current_user.id
    ).order_by(CopilotChat.actualizado_en.desc()).limit(MAX_CHAT_LIST).all()
    return [{"id": c.id, "titulo": c.titulo, "actualizado_en": c.actualizado_en} for c in chats]

@router.post("/chats")
def create_chat(db: Session = Depends(get_db), current_user: Usuario = Depends(get_current_user)):
    new_chat = CopilotChat(usuario_id=current_user.id, titulo="Nuevo Chat")
    db.add(new_chat)
    db.commit()
    db.refresh(new_chat)
    return {"id": new_chat.id, "titulo": new_chat.titulo, "actualizado_en": new_chat.actualizado_en}

@router.get("/chats/{chat_id}")
def get_chat_history(chat_id: int, db: Session = Depends(get_db), current_user: Usuario = Depends(get_current_user)):
    # El historial es persistencia de la aplicacion, no una consulta analitica
    # generada por la IA. Leerlo con la sesion normal evita depender de grants
    # RO sobre tablas de chat creadas despues de configurar DATABASE_RO_URL.
    chat = db.query(CopilotChat).filter(CopilotChat.id == chat_id, CopilotChat.usuario_id == current_user.id).first()
    if not chat:
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail="Chat no encontrado")
    mensajes = db.query(CopilotMessage).filter(
        CopilotMessage.chat_id == chat_id
    ).order_by(CopilotMessage.creado_en.desc()).limit(MAX_CHAT_HISTORY_MESSAGES).all()
    mensajes.reverse()
    return [
        {"id": m.id, "role": m.rol, "content": m.contenido, "creado_en": _fecha_mensaje_api(m.creado_en)}
        for m in mensajes
    ]

@router.get("/chats/{chat_id}/export")
def export_chat(
    chat_id: int,
    format: Literal["csv", "xlsx", "pdf"] = "pdf",
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
):
    """Exporta el historial visible de un chat, sin marcadores internos de IA."""
    import csv
    import html
    import io
    from fastapi.responses import StreamingResponse

    chat = db.query(CopilotChat).filter(
        CopilotChat.id == chat_id,
        CopilotChat.usuario_id == current_user.id,
    ).first()
    if not chat:
        raise HTTPException(status_code=404, detail="Chat no encontrado")

    mensajes = db.query(CopilotMessage).filter(
        CopilotMessage.chat_id == chat.id,
    ).order_by(CopilotMessage.creado_en.asc()).all()
    if not mensajes:
        raise HTTPException(status_code=400, detail="El chat no tiene mensajes para exportar")

    def limpiar_contenido(contenido: str) -> str:
        return re.sub(r"<!-- (?:sql_export|copilot_metrics|copilot_followups|copilot_chart|copilot_report): [\s\S]*? -->", "", contenido).strip()

    filas = [
        {
            "Fecha": mensaje.creado_en.strftime("%H:%M %d-%m-%Y"),
            "Rol": "Usuario" if mensaje.rol == "user" else "AI Copilot",
            "Mensaje": limpiar_contenido(mensaje.contenido),
        }
        for mensaje in mensajes
    ]
    nombre_base = f"conversacion_copilot_{chat.id}"

    if format == "csv":
        output = io.StringIO()
        output.write("\ufeff")
        writer = csv.DictWriter(output, fieldnames=("Fecha", "Rol", "Mensaje"))
        writer.writeheader()
        writer.writerows(filas)
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f"attachment; filename={nombre_base}.csv"},
        )

    if format == "xlsx":
        import pandas as pd

        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame(filas).to_excel(writer, index=False, sheet_name="Conversación")
            worksheet = writer.book["Conversación"]
            worksheet.column_dimensions["A"].width = 20
            worksheet.column_dimensions["B"].width = 16
            worksheet.column_dimensions["C"].width = 100
            for cell in worksheet[1]:
                cell.font = __import__("openpyxl").styles.Font(bold=True)
            for row in worksheet.iter_rows(min_row=2, min_col=3, max_col=3):
                row[0].alignment = __import__("openpyxl").styles.Alignment(wrap_text=True, vertical="top")
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={nombre_base}.xlsx"},
        )

    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, rightMargin=1.6 * cm, leftMargin=1.6 * cm, topMargin=1.6 * cm, bottomMargin=1.6 * cm)
    styles = getSampleStyleSheet()
    titulo = ParagraphStyle("CopilotTitle", parent=styles["Title"], textColor=HexColor("#0F172A"), fontSize=18, leading=22)
    meta = ParagraphStyle("CopilotMeta", parent=styles["Normal"], textColor=HexColor("#64748B"), fontSize=8, leading=11)
    usuario = ParagraphStyle("CopilotUser", parent=styles["BodyText"], backColor=HexColor("#DBEAFE"), textColor=HexColor("#172554"), borderPadding=8, leading=14, spaceAfter=4)
    asistente = ParagraphStyle("CopilotAssistant", parent=styles["BodyText"], backColor=HexColor("#F1F5F9"), textColor=HexColor("#0F172A"), borderPadding=8, leading=14, spaceAfter=4)
    historia = [Paragraph(html.escape(chat.titulo), titulo), Spacer(1, 0.25 * cm)]
    for fila in filas:
        historia.append(Paragraph(f"{fila['Rol']} · {fila['Fecha']}", meta))
        texto = html.escape(fila["Mensaje"]).replace("\n", "<br/>")
        historia.append(Paragraph(texto or "—", usuario if fila["Rol"] == "Usuario" else asistente))
        historia.append(Spacer(1, 0.18 * cm))
    document.build(historia)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={nombre_base}.pdf"},
    )

@router.put("/chats/{chat_id}")
def rename_chat(chat_id: int, payload: RenameChatRequest, db: Session = Depends(get_db), current_user: Usuario = Depends(get_current_user)):
    chat = db.query(CopilotChat).filter(
        CopilotChat.id == chat_id,
        CopilotChat.usuario_id == current_user.id,
    ).first()
    if not chat:
        raise HTTPException(status_code=404, detail="Chat no encontrado")
    chat.titulo = payload.titulo.strip()
    chat.actualizado_en = datetime.utcnow()
    db.commit()
    db.refresh(chat)
    return {"id": chat.id, "titulo": chat.titulo, "actualizado_en": chat.actualizado_en}

@router.delete("/chats/{chat_id}")
def delete_chat(chat_id: int, db: Session = Depends(get_db), current_user: Usuario = Depends(get_current_user)):
    chat = db.query(CopilotChat).filter(CopilotChat.id == chat_id, CopilotChat.usuario_id == current_user.id).first()
    if chat:
        db.delete(chat)
        db.commit()
    return {"success": True}

@router.post("/chat")
@limiter.limit("5/minute")
def copilot_chat(request: Request, payload: ChatRequest, db: Session = Depends(get_db), current_user: Usuario = Depends(get_current_user)):
    try:
        # Si no hay chat_id, se crea uno
        if not payload.chat_id:
            chat = CopilotChat(usuario_id=current_user.id, titulo=payload.history[-1].content[:30] + "...")
            db.add(chat)
            db.commit()
            db.refresh(chat)
            chat_id = chat.id
        else:
            chat_id = payload.chat_id
            chat = db.query(CopilotChat).filter(CopilotChat.id == chat_id, CopilotChat.usuario_id == current_user.id).first()
            if not chat:
                raise HTTPException(status_code=404, detail="Chat no encontrado")
            chat.actualizado_en = datetime.utcnow()
            db.commit()

        # Extraer el historial para procesar (formato para process_copilot_chat)
        # Podríamos usar el de la BD, pero process_copilot_chat requiere history de dicts
        # Reconstruimos la historia completa desde la BBDD + el nuevo mensaje
        mensajes_previos = db.query(CopilotMessage).filter(
            CopilotMessage.chat_id == chat_id
        ).order_by(CopilotMessage.creado_en.desc()).limit(MAX_CHAT_CONTEXT_MESSAGES).all()
        mensajes_previos.reverse()
        
        # El payload.history normalmente trae solo el nuevo mensaje si el frontend está bien diseñado, 
        # pero si envía todo el historial, guardamos solo el último para evitar duplicados.
        nuevo_mensaje = payload.history[-1]
        
        # Guardar la pregunta del usuario
        user_msg_db = CopilotMessage(chat_id=chat_id, rol=nuevo_mensaje.role, contenido=nuevo_mensaje.content)
        db.add(user_msg_db)
        db.commit()

        # Reconstruir history para la IA
        history_dicts = [{"role": m.rol, "content": m.contenido} for m in mensajes_previos]
        history_dicts.append({"role": nuevo_mensaje.role, "content": nuevo_mensaje.content})

        # Obtener contexto de negocio
        empresa = db.query(models.Empresa).filter(models.Empresa.id == current_user.empresa_id).first()
        contexto_negocio = empresa.contexto_negocio if empresa else ""

        if payload.libreria_doc_ids:
            try:
                from ..models import LibreriaDocumento
                libreria_docs = db.query(LibreriaDocumento).filter(
                    LibreriaDocumento.empresa_id == current_user.empresa_id,
                    LibreriaDocumento.id.in_(payload.libreria_doc_ids)
                ).all()
                if libreria_docs:
                    lib_ctx = "\n\n=== DOCUMENTOS DE REFERENCIA (LibrerIA) ===\nUsa estos documentos como base de conocimiento, terminologia y contexto del negocio:\n\n"
                    for doc in libreria_docs:
                        lib_ctx += f"--- {doc.filename} ({doc.department}) ---\n{doc.content_text[:3000]}\n\n"
                    contexto_negocio = (contexto_negocio or "") + lib_ctx
            except Exception as e_lib:
                logger.warning(f"Error cargando docs LibrerIA: {e_lib}")

        # DB RO para queries analíticas de la IA
        try:
            contexto_intencion = json.loads(chat.contexto_analitico or "{}")
        except (TypeError, json.JSONDecodeError):
            logger.warning("[COPILOT] Contexto analítico inválido para chat_id=%s; se reinicia.", chat_id)
            contexto_intencion = {}
        from ..database import SessionLocalRO
        db_ro = SessionLocalRO()
        try:
            reply = process_copilot_chat(
                db=db_ro, 
                history=history_dicts, 
                empresa_id=current_user.empresa_id,
                model_preference=payload.model_preference,
                contexto=contexto_negocio,
                contexto_intencion=contexto_intencion,
            )
        finally:
            db_ro.close()

        # El siguiente turno reutiliza este contrato validado, no la prosa de
        # la respuesta. Si la petición no era analítica se conserva el último
        # contrato para que «lo mismo» siga siendo interpretable.
        from ..copilot_orchestrator import analizar_intencion, serializar_intento
        intento_actual, _ = analizar_intencion(history_dicts, contexto_intencion)
        if intento_actual:
            chat.contexto_analitico = json.dumps(serializar_intento(intento_actual), ensure_ascii=False)
            
        # Guardar respuesta IA
        assistant_msg_db = CopilotMessage(chat_id=chat_id, rol="assistant", contenido=reply)
        db.add(assistant_msg_db)
        db.commit()

        return {"reply": reply, "chat_id": chat_id, "message_id": assistant_msg_db.id}
    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        logger.error(f"[COPILOT CHAT] Error procesando chat: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail="Error interno procesando el chat.")

@router.get("/chat/message/{message_id}/export")
def download_message_export(
    message_id: int,
    format: str = "csv",
    current_user: Usuario = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    import io
    import csv
    from fastapi.responses import StreamingResponse
    from ..database import SessionLocalRO
    from ..copilot_service import execute_sql
    
    # 1. Recuperar el mensaje asegurando que pertenece a un chat del usuario
    msg = db.query(CopilotMessage).join(CopilotChat).filter(
        CopilotMessage.id == message_id,
        CopilotChat.usuario_id == current_user.id
    ).first()
    
    if not msg:
        raise HTTPException(status_code=404, detail="Mensaje no encontrado")

    informe = extract_signed_report_artifact(msg.contenido)
    if informe:
        return _exportar_informe_ejecutivo(informe, message_id, format)
        
    # 2. Extraer el SQL firmado
    signed_export = extract_signed_sql_export(msg.contenido)
    if not signed_export:
        raise HTTPException(status_code=400, detail="Este mensaje no contiene datos exportables.")
    sql_query, trusted_query, query_params = signed_export
    parametros_requeridos = set(re.findall(r":([A-Za-z_][A-Za-z0-9_]*)", sql_query)) - {"empresa_id"}
    if parametros_requeridos and not query_params:
        raise HTTPException(
            status_code=409,
            detail="Esta respuesta se generó antes de guardar los parámetros de exportación. Repite la consulta para descargarla.",
        )
        
    # 3. Re-ejecutar SQL en read-only
    db_ro = SessionLocalRO()
    try:
        raw_data, error = execute_sql(
            db_ro,
            sql_query,
            current_user.empresa_id,
            query_params,
            trusted_query=trusted_query,
        )
        if error or not raw_data:
            raise HTTPException(status_code=500, detail="Error ejecutando la consulta original o sin datos.")
            
        if format == "xlsx":
            import pandas as pd
            
            output = io.BytesIO()
            if len(raw_data) > 0:
                df = pd.DataFrame(raw_data)
                df.to_excel(output, index=False, engine='openpyxl')
            
            output.seek(0)
            return StreamingResponse(
                output, 
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 
                headers={"Content-Disposition": f"attachment; filename=exportacion_ia_{message_id}.xlsx"}
            )
        else:
            # Default to CSV
            output = io.StringIO()
            output.write('\ufeff') # BOM (Byte Order Mark) para que Excel lea UTF-8 correctamente (tildes/ñ)
            if len(raw_data) > 0:
                writer = csv.DictWriter(output, fieldnames=raw_data[0].keys())
                writer.writeheader()
                for row in raw_data:
                    writer.writerow(row)
                    
            output.seek(0)
            return StreamingResponse(
                output, 
                media_type="text/csv", 
                headers={"Content-Disposition": f"attachment; filename=exportacion_ia_{message_id}.csv"}
            )
    finally:
        db_ro.close()
